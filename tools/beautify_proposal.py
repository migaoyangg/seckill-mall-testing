from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor


SOURCE = Path(
    "/Users/migaoyang/Library/Containers/com.tencent.xinWeChat/Data/Documents/"
    "xwechat_files/wxid_dnzdvnx9nekd22_83db/msg/file/2026-07/"
    "华图成长罗盘-开题报告-扩充版-美化.docx"
)
OUTPUT = Path("deliverables/华图成长罗盘-开题报告-排版美化版.docx")

NAVY = "173F6B"
BLUE = "2F75B5"
LIGHT_BLUE = "F3F7FB"
TEXT = RGBColor(40, 54, 70)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_borders(cell, color="D9E2F3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "7")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "打开文档后右键此处，选择“更新域”生成目录。"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])


def set_update_fields(document):
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def restore_markdown_bold(paragraph):
    text = paragraph.text
    if "**" not in text:
        return
    first = paragraph.runs[0] if paragraph.runs else None
    base_name = first.font.name if first and first.font.name else "微软雅黑"
    base_size = first.font.size if first and first.font.size else Pt(10.5)
    base_bold = first.bold if first else None
    paragraph.clear()
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run(part[2:-2] if bold else part)
        run.font.name = base_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), base_name)
        run.font.size = base_size
        run.bold = True if bold else base_bold


def add_hyperlink(paragraph, url, text, template_run=None):
    """Append a Word-native external hyperlink while preserving the base font."""
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    if template_run is not None and template_run._r.rPr is not None:
        run.append(deepcopy(template_run._r.rPr))
    r_pr = run.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        run.insert(0, r_pr)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def linkify_urls(paragraph):
    """Turn displayed HTTP(S) addresses into clickable DOCX hyperlinks."""
    text = paragraph.text
    url_pattern = re.compile(r"https?://[^\s<>\"\u3002\uff0c\uff1b\u3001\uff09\u3011]+")
    matches = list(url_pattern.finditer(text))
    if not matches:
        return
    template = paragraph.runs[0] if paragraph.runs else None
    paragraph.clear()
    cursor = 0
    for match in matches:
        if match.start() > cursor:
            prefix = paragraph.add_run(text[cursor:match.start()])
            if template is not None and template._r.rPr is not None:
                prefix._r.insert(0, deepcopy(template._r.rPr))
        raw_url = match.group(0)
        # Keep sentence punctuation outside the clickable address.
        url = raw_url.rstrip(".,;:!?)】")
        add_hyperlink(paragraph, url, url, template)
        if len(url) < len(raw_url):
            suffix = paragraph.add_run(raw_url[len(url):])
            if template is not None and template._r.rPr is not None:
                suffix._r.insert(0, deepcopy(template._r.rPr))
        cursor = match.end()
    if cursor < len(text):
        suffix = paragraph.add_run(text[cursor:])
        if template is not None and template._r.rPr is not None:
            suffix._r.insert(0, deepcopy(template._r.rPr))


def all_paragraphs(document):
    for p in document.paragraphs:
        yield p
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def heading_level(text):
    if re.match(r"^[一二三四五六七八九十]+、", text):
        return 1
    if re.match(r"^\d+\.\d+", text):
        return 2
    return 0


def make_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    h1 = document.styles["Heading 1"]
    h1.font.name = "微软雅黑"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(23, 63, 107)
    h1.paragraph_format.space_before = Pt(24)
    h1.paragraph_format.space_after = Pt(12)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.page_break_before = True

    h2 = document.styles["Heading 2"]
    h2.font.name = "微软雅黑"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(47, 117, 181)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(8)
    h2.paragraph_format.keep_with_next = True

    toc_title = document.styles.add_style("目录标题", WD_STYLE_TYPE.PARAGRAPH)
    toc_title.base_style = document.styles["Heading 1"]
    toc_title.font.color.rgb = RGBColor(23, 63, 107)
    toc_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.page_break_before = False


def build_document():
    document = Document(SOURCE)
    make_styles(document)
    set_update_fields(document)

    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.18)
        section.right_margin = Cm(3.18)
        section.different_first_page_header_footer = True
        section.first_page_header.paragraphs[0].text = ""
        section.first_page_footer.paragraphs[0].text = ""

    # Promote the existing visual hierarchy to real Word headings.
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        level = heading_level(text)
        if level == 1:
            paragraph.style = document.styles["Heading 1"]
            add_bottom_border(paragraph)
        elif level == 2:
            paragraph.style = document.styles["Heading 2"]
        elif text:
            paragraph.style = document.styles["Normal"]
            if text.startswith("•"):
                paragraph.paragraph_format.left_indent = Cm(0.74)
                paragraph.paragraph_format.first_line_indent = Cm(-0.37)

    # Decorative unicode rules are replaced by real heading borders.
    for paragraph in list(document.paragraphs):
        if paragraph.text.strip() and set(paragraph.text.strip()) == {"━"}:
            paragraph._element.getparent().remove(paragraph._element)

    # The source uses empty, manual page-break paragraphs before every chapter.
    # Heading 1 now owns pagination, which avoids duplicated blank pages.
    for paragraph in list(document.paragraphs):
        if not paragraph.text.strip() and 'w:type="page"' in paragraph._p.xml:
            paragraph._element.getparent().remove(paragraph._element)

    for paragraph in all_paragraphs(document):
        restore_markdown_bold(paragraph)
        linkify_urls(paragraph)

    for table_index, table in enumerate(document.tables):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        # A single-cell table after the cover is a narrative callout, not a table
        # header. Treat it as a readable note instead of a dense dark-blue block.
        is_callout = table_index > 0 and len(table.rows) == 1 and len(table.columns) == 1
        if not is_callout:
            mark_header_row(table.rows[0])
        for row_index, row in enumerate(table.rows):
            row.height = None
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_borders(cell)
                if row_index == 0 and not is_callout:
                    set_cell_shading(cell, NAVY)
                elif is_callout:
                    set_cell_shading(cell, LIGHT_BLUE)
                elif row_index % 2 == 0:
                    set_cell_shading(cell, LIGHT_BLUE)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.15
                    for run in paragraph.runs:
                        run.font.name = "微软雅黑"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                        run.font.size = Pt(9.5)
                        if row_index == 0 and not is_callout:
                            run.font.color.rgb = RGBColor(255, 255, 255)
                            run.bold = True
                    if row_index == 0 and not is_callout:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Make the cover stand alone and insert a dynamic table of contents after it.
    cover_anchor = next(
        p for p in document.paragraphs if "企业命题：华图教育" in p.text
    )
    page_break = OxmlElement("w:p")
    page_break_pr = OxmlElement("w:pPr")
    page_break.append(page_break_pr)
    page_break_r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    page_break_r.append(br)
    page_break.append(page_break_r)

    toc_title = OxmlElement("w:p")
    toc_title_pr = OxmlElement("w:pPr")
    toc_style = OxmlElement("w:pStyle")
    toc_style.set(qn("w:val"), "目录标题")
    toc_title_pr.append(toc_style)
    toc_title.append(toc_title_pr)
    title_run = OxmlElement("w:r")
    title_text = OxmlElement("w:t")
    title_text.text = "目录"
    title_run.append(title_text)
    toc_title.append(title_run)

    toc_para = OxmlElement("w:p")
    toc_para_pr = OxmlElement("w:pPr")
    toc_para.append(toc_para_pr)
    toc_para_r = OxmlElement("w:r")
    toc_para.append(toc_para_r)
    # Create fields by way of a temporary Paragraph proxy after insertion.
    anchor = cover_anchor._p
    anchor.addnext(toc_para)
    anchor.addnext(toc_title)
    anchor.addnext(page_break)
    # The paragraph proxy is only needed to add the TOC field nodes cleanly.
    from docx.text.paragraph import Paragraph
    add_toc_field(Paragraph(toc_para, document._body))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build_document()
