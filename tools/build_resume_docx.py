from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "/Users/migaoyang/Projects/seckill-mall/deliverables/测试开发工程师简历_一页版.docx"
NAVY = RGBColor(31, 78, 121)
DARK = RGBColor(35, 35, 35)
GRAY = RGBColor(95, 95, 95)


def set_cn_font(run, name="Heiti SC", size=9.8, bold=False, color=DARK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_cell_margin(cell, top=20, start=25, bottom=20, end=25):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def format_para(p, before=0, after=0, line=1.03):
    f = p.paragraph_format
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = line


def add_text(p, text, bold=False, size=9.8, color=DARK):
    r = p.add_run(text)
    set_cn_font(r, size=size, bold=bold, color=color)
    return r


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    format_para(p, before=4.5, after=2.5, line=1)
    add_text(p, text, bold=True, size=11.8, color=NAVY)


def bullet(doc, text, marker=None):
    p = doc.add_paragraph(style="List Bullet")
    format_para(p, after=1.4, line=1.1)
    p.paragraph_format.left_indent = Inches(0.17)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    add_text(p, text, size=9.35)
    if marker:
        add_text(p, marker, size=9.0, color=GRAY)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.34)
sec.bottom_margin = Inches(0.32)
sec.left_margin = Inches(0.48)
sec.right_margin = Inches(0.48)
sec.header_distance = Inches(0.15)
sec.footer_distance = Inches(0.15)

normal = doc.styles["Normal"]
normal.font.name = "Heiti SC"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Heiti SC")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Heiti SC")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
normal.font.size = Pt(9.8)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(0)

for style_name in ("Title", "Heading 1", "Heading 2"):
    style = doc.styles[style_name]
    style.font.name = "Heiti SC"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Heiti SC")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Heiti SC")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
    style.font.color.rgb = RGBColor(0, 0, 0)

# Header with photo placeholder
head = doc.add_table(rows=1, cols=2)
head.autofit = False
head.columns[0].width = Inches(6.55)
head.columns[1].width = Inches(0.95)
left = head.cell(0, 0)
right = head.cell(0, 1)
for c in (left, right):
    set_cell_margin(c, 15, 15, 15, 15)
    set_cell_border(c, color="FFFFFF", size="0")
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

p = left.paragraphs[0]
p.style = doc.styles["Title"]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
format_para(p, after=1.5, line=1)
add_text(p, "蒙冠源", bold=True, size=20, color=RGBColor(0, 0, 0))

p = left.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
format_para(p, after=2, line=1)
add_text(p, "求职意向  软件测试开发工程师", bold=True, size=10.2, color=NAVY)

p = left.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
format_para(p, after=1, line=1)
add_text(p, "电话  16677266727   ｜   邮箱  migoyang66@163.com   ｜   意向城市  北京 / 天津", size=8.7, color=GRAY)

set_cell_border(right, color="BFBFBF", size="8")
shade_cell(right, "F2F2F2")
p = right.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
format_para(p, before=24, after=24, line=1)
add_text(p, "证件照\n位置", bold=True, size=9.0, color=GRAY)

section_heading(doc, "教育经历")
t = doc.add_table(rows=1, cols=2)
t.autofit = False
t.columns[0].width = Inches(5.9)
t.columns[1].width = Inches(1.55)
for c in t.rows[0].cells:
    set_cell_margin(c, 0, 0, 0, 0)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(c, color="FFFFFF", size="0")
p = t.cell(0, 0).paragraphs[0]
format_para(p, line=1)
add_text(p, "天津中德应用技术大学  ｜  软件工程（全日制本科）", bold=True, size=9.3)
p = t.cell(0, 1).paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
format_para(p, line=1)
add_text(p, "2023.09 - 2027.06", size=8.8, color=GRAY)

p = doc.add_paragraph()
format_para(p, before=1, after=1.5, line=1.04)
add_text(p, "主修课程  ", bold=True, size=8.8, color=NAVY)
add_text(p, "软件测试基础、数据结构与算法、操作系统、计算机网络、数据库原理、Python 程序设计、Java 程序设计", size=8.75, color=GRAY)

section_heading(doc, "项目经历")

def project_header(title, date, tech, link):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(5.95)
    table.columns[1].width = Inches(1.55)
    for cell in table.rows[0].cells:
        set_cell_margin(cell, 0, 0, 0, 0)
        set_cell_border(cell, color="FFFFFF", size="0")
    p1 = table.cell(0, 0).paragraphs[0]
    format_para(p1, before=2, after=1, line=1)
    add_text(p1, title, bold=True, size=10.7, color=RGBColor(0, 0, 0))
    add_text(p1, "   GitHub  ", bold=True, size=8.0, color=GRAY)
    add_text(p1, link, size=8.0, color=GRAY)
    p2 = table.cell(0, 1).paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    format_para(p2, before=1, after=0.5, line=1)
    add_text(p2, date, size=9.0, color=GRAY)
    p = doc.add_paragraph()
    format_para(p, after=1.3, line=1)
    add_text(p, tech, bold=True, size=8.65, color=NAVY)


def project_description(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.10)
    p.paragraph_format.right_indent = Inches(0.04)
    format_para(p, after=1.8, line=1.06)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left_border = OxmlElement("w:left")
    left_border.set(qn("w:val"), "single")
    left_border.set(qn("w:sz"), "18")
    left_border.set(qn("w:space"), "6")
    left_border.set(qn("w:color"), "2F75B5")
    pBdr.append(left_border)
    pPr.append(pBdr)
    add_text(p, "项目描述  ", bold=True, size=8.9, color=NAVY)
    add_text(p, text, size=8.95, color=GRAY)


def bullet_labeled(label, text, marker=None):
    p = doc.add_paragraph(style="List Bullet")
    format_para(p, after=1.2, line=1.08)
    p.paragraph_format.left_indent = Inches(0.17)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    add_text(p, label + "：", bold=True, size=9.25)
    add_text(p, text, size=9.2)
    if marker:
        add_text(p, "  " + marker, size=8.9, color=GRAY)


project_header(
    "秒购商城服务端测试自动化框架",
    "2026.07 至今",
    "Spring Boot 3  ·  MySQL  ·  Redis  ·  JUnit 5  ·  Mockito  ·  MockMvc  ·  pytest  ·  JMeter",
    "【项目地址】",
)
project_description(
    "针对秒杀商城在高并发场景下的鉴权、订单状态、支付幂等、库存扣减和异步下单风险，建设覆盖单元、Web、接口、数据一致性及性能场景的自动化测试框架。",
)
bullet_labeled("分层测试体系", "基于 JUnit 5、Mockito、MockMvc 建立服务层、Web 层与调度任务测试，覆盖鉴权、订单状态、支付幂等、库存回滚及分布式锁，38 条用例全部通过。")
bullet_labeled("接口自动化框架", "使用 pytest + requests 封装配置、Token、Session、公共断言和动态数据工厂，覆盖用户、商品、订单、退款及秒杀异步链路，28 条真实 HTTP 用例全部通过。")
bullet_labeled("鉴权与异常治理", "定位参数类型错误返回 HTTP 500、商品管理路径漏拦截及 Token fixture 相互覆盖问题，补充异常处理、权限配置和函数级数据隔离并完成回归。")
bullet_labeled("CI 工程化", "编写 GitHub Actions 与 Docker Compose 流程，覆盖 MySQL、Redis、RabbitMQ、应用健康检查、真实 HTTP 回归、JUnit/HTML 报告与服务日志归档。")
bullet_labeled("性能与一致性", "使用 JMeter 对秒杀入口执行 100 用户瞬时并发，吞吐 123.30 请求/秒、P95 255.6 ms、错误率 0%；100 笔订单异步落库，重复订单、库存差值及队列积压均为 0。")

project_header(
    "秒购商城 Android 客户端自动化测试框架",
    "2026.09 至今",
    "Python  ·  pytest  ·  Appium  ·  UiAutomator2  ·  ADB  ·  Page Object  ·  Allure｜被测端：Kotlin + Retrofit",
    "【项目地址】",
)
project_description(
    "以秒购商城 Android 客户端为被测对象，打通 Android Emulator、Appium/UiAutomator2、客户端、Spring Boot API 与 MySQL/Redis 的真实端到端测试链路。",
)
bullet_labeled("页面对象架构", "设计 BasePage、LoginPage、HomePage、ProductDialog 4 个页面对象类，分离元素定位、页面操作与断言，封装显式等待、点击输入、文本读取和截图能力。")
bullet_labeled("配置与数据保护", "使用 fixture 管理 Driver、登录前置与测试数据；UI 下单并验证后通过业务 API 自动取消新增订单、恢复库存，账号与设备配置均由环境变量注入。")
bullet_labeled("端到端回归", "编写 13 条 UI 自动化用例，覆盖登录校验、应用重启状态恢复、商品详情、页面切换、订单查询与创建订单，全量回归 13 条全部通过，耗时 202.47 秒。")
bullet_labeled("失败取证与稳定性", "失败时自动保存截图、页面 XML 和 Logcat 并生成 HTML/Allure 结果；定位弹窗资源 ID 与异步文案读取问题，修复后复测通过。")
bullet_labeled("Android CI", "编写 GitHub Actions 工作流，自动启动后端依赖与 Android Emulator、构建 APK、启动 Appium、执行 smoke 用例并归档测试证据。")

section_heading(doc, "专业技能")
skill_texts = [
    ("编程与自动化", "Java、Python、SQL；JUnit 5、Mockito、MockMvc、pytest、requests、Appium、Page Object"),
    ("接口与性能", "Postman、curl、JMeter；掌握接口关联、参数化及 P95/P99、错误率等指标"),
    ("数据与工程", "MySQL、Redis、RabbitMQ；Git、Maven、Linux、ADB、Docker Compose、GitHub Actions；了解 Jenkins"),
    ("AI 工具", "使用 ChatGPT、Codex 辅助用例设计、脚本开发、代码检查、问题排查和文档整理"),
]
for label, value in skill_texts:
    p = doc.add_paragraph(style="List Bullet")
    format_para(p, after=1.0, line=1.04)
    p.paragraph_format.left_indent = Inches(0.17)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    add_text(p, label + "：", bold=True, size=9.0)
    add_text(p, value, size=8.95)

doc.save(OUT)
print(OUT)
